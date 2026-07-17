"""Build the plain, editable Word report for Formative 2."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "reports" / "figures"
OUTPUT = ROOT / "output" / "word" / "formative2_multimodal_preprocessing_report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value
        set_cell_shading(cell, "E7E6E6")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9.5)
        if widths:
            cell.width = widths[index]
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            if widths:
                cells[index].width = widths[index]
    document.add_paragraph()
    return table


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def add_figure(document: Document, filename: str, caption: str, width: float = 6.3) -> None:
    path = FIGURES / filename
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(document, caption)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def build_report() -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(10)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.space_before = Pt(7)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)

    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Formative 2: Multimodal Data Preprocessing")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("User Identity and Product Recommendation System").bold = True
    group = document.add_paragraph()
    group.alignment = WD_ALIGN_PARAGRAPH.CENTER
    group.add_run("Group 19")
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run("Prepared by Christian Tonny Gentil Iradukunda and Group 19")
    document.add_paragraph()

    document.add_heading("1. Introduction", level=1)
    document.add_paragraph(
        "This project simulates a secure product recommendation system. A face image is checked first. "
        "If it is accepted, the product model runs but does not show its answer yet. A voice sample is then "
        "checked. The recommendation is only displayed when the face and voice identities match. Failed "
        "checks return Access Denied."
    )
    document.add_paragraph(
        "The project combines two tabular datasets with facial images and short audio recordings. The work "
        "includes data cleaning, merge validation, exploratory data analysis, feature engineering, media "
        "augmentation, three Random Forest models, evaluation, and a command-line demonstration."
    )

    document.add_heading("2. Tabular Data Preprocessing", level=1)
    document.add_paragraph(
        "The customer social profile dataset originally contained 155 rows and 5 exact duplicates. After "
        "duplicate removal, 150 social observations remained for 84 unique customers. Customer IDs such as "
        "A178 were converted to numeric values so they could match the transaction data."
    )
    document.add_paragraph(
        "The transaction dataset contained 150 rows. Ten customer ratings were missing and were filled using "
        "the median value of 3.0. Purchase dates were converted to datetime values. Transaction IDs and product "
        "categories did not have missing values."
    )
    document.add_paragraph(
        "Some customers had several social observations. A direct join would duplicate their transactions, "
        "so social data was summarized to one row per customer first. The engineered features include mean and "
        "maximum engagement, mean and maximum purchase interest, social profile count, platform count, strongest "
        "platform, and modal review sentiment."
    )
    document.add_paragraph(
        "A validated many-to-one inner merge produced 117 transactions. Thirty-three transaction rows were "
        "excluded because those customers had no matching social profile. Post-merge checks confirmed that no "
        "transaction ID was duplicated. Purchase month, day of week, and quarter were also extracted. Customer "
        "rating was excluded from the model inputs to avoid using information normally collected after purchase."
    )
    add_table(
        document,
        ["Validation item", "Result"],
        [
            ["Raw social rows", "155"],
            ["Exact social duplicates removed", "5"],
            ["Clean social observations / customers", "150 / 84"],
            ["Raw transaction rows", "150"],
            ["Missing ratings imputed", "10 using median 3.0"],
            ["Merged rows / unmatched transactions", "117 / 33"],
            ["Duplicated transaction IDs after merge", "0"],
        ],
    )

    document.add_heading("3. Exploratory Data Analysis", level=1)
    document.add_paragraph(
        "The merged target was fairly balanced: Sports had 28 rows, Electronics 27, Clothing 22, Groceries "
        "20, and Books 20. Purchase amount ranged from 62 to 495, with a mean of 287.91. Mean engagement was "
        "74.15 and mean purchase interest was 3.08."
    )
    document.add_paragraph(
        "The distribution, boxplot, and correlation plots show a lot of overlap between product categories. "
        "Numeric correlations are weak. This suggests that the available features will not strongly predict "
        "which product category a customer purchases."
    )
    add_figure(document, "tabular_eda.png", "Figure 1. Target distribution, purchase amount outliers, and numeric correlations.")

    document.add_heading("4. Image Data Processing", level=1)
    document.add_paragraph(
        "Images were organized in one folder per group member. Each image was converted to RGB and fitted to "
        "a consistent size. Three variants were created from every original: a 15-degree rotation, horizontal "
        "flip, and grayscale conversion."
    )
    document.add_paragraph(
        "The feature file contains normalized red, green, blue, and grayscale histograms; mean and standard "
        "deviation of brightness; and a 16 by 16 grayscale pixel embedding. This gives 322 numeric image "
        "features per row, plus metadata columns."
    )
    document.add_paragraph(
        "The current local dataset had three Christian images and one Yvette image. Four originals produced "
        "16 rows after adding the original and three variants. Augmented copies from the same original were "
        "kept on the same side of the train/test split to reduce leakage."
    )
    add_figure(document, "christian_image_samples.png", "Figure 2. Christian's neutral, smiling, and surprised image samples.")

    document.add_heading("5. Audio Data Processing", level=1)
    document.add_paragraph(
        "Phone recordings were decoded as mono audio at 16 kHz. Every original received three augmentations: "
        "a two-semitone pitch shift, a 0.90 time stretch, and low-level background noise."
    )
    document.add_paragraph(
        "Audio features were the mean and standard deviation of 13 MFCC coefficients, spectral roll-off, RMS "
        "energy, zero-crossing rate, and duration. This gives 30 numeric features per row, plus metadata columns."
    )
    document.add_paragraph(
        "Christian and Mahlet each had two original recordings. Four originals produced 16 rows after augmentation. "
        "The waveforms show clear speech regions, while the spectrograms show that most speech energy is in the "
        "lower frequencies."
    )
    add_figure(document, "audio_waveforms_spectrograms.png", "Figure 3. Waveforms and spectrograms for the available group members.")

    document.add_heading("6. Model Creation and Evaluation", level=1)
    document.add_paragraph(
        "Three Random Forest classifiers were implemented. The product pipeline median-imputes numeric inputs "
        "and one-hot encodes categorical inputs. The face and voice models use original-file grouping so an "
        "augmented version of a test sample is not placed in training."
    )
    add_table(
        document,
        ["Model", "Accuracy", "Macro F1", "Log loss"],
        [
            ["Product recommendation", "0.167", "0.177", "1.746"],
            ["Facial recognition", "1.000", "1.000", "0.088"],
            ["Voiceprint verification", "1.000", "1.000", "0.096"],
        ],
    )
    document.add_paragraph(
        "The product model was evaluated on 30 held-out rows. Its low result agrees with the weak relationships "
        "found in EDA, so it is functional but not production quality."
    )
    document.add_paragraph(
        "The voice evaluation included Christian and Mahlet. The face evaluation only included Christian because "
        "Yvette had one original image. Therefore, the perfect face result is incomplete and should not be "
        "generalized. Both authentication datasets are very small, which can make the scores unstable."
    )
    add_figure(document, "product_confusion_matrix.png", "Figure 4. Product recommendation confusion matrix.", width=4.7)

    document.add_heading("7. Multimodal System Demonstration", level=1)
    document.add_paragraph("The command-line application uses this sequence:")
    numbered = [
        "Read the face image and predict its identity.",
        "Reject the sample if face confidence is below 85%.",
        "Run the product model, but keep the answer hidden.",
        "Read the voice sample and predict its identity.",
        "Reject the sample if voice confidence is too low.",
        "Compare the face identity with the voice identity.",
        "Display the product only when both identities match.",
    ]
    for item in numbered:
        document.add_paragraph(item, style="List Number")
    document.add_paragraph(
        "In the successful test, Christian's neutral face scored 89.7% confidence and his Yes, approve audio "
        "scored 99.3%. The identities matched, access was approved, and the product model predicted Clothing "
        "for the sample transaction."
    )
    document.add_paragraph(
        "Two denial paths were tested. A non-face image scored 84.0%, which is below the 85% face threshold, "
        "so access was denied before showing a recommendation. Christian's face was also combined with Mahlet's "
        "voice. Both samples were recognized, but the identities did not match, so access was denied."
    )

    document.add_heading("8. Group Contributions", level=1)
    add_table(
        document,
        ["Member", "Contribution"],
        [
            ["Christian Tonny Gentil Iradukunda", "Submitted three expressions and two voice phrases; completed cleaning, merge logic, feature engineering, model integration, testing, CLI simulation, notebook, and report."],
            ["Hassan", "Created and shared the group data folders and coordinated member uploads."],
            ["Mahlet Tilahun", "Submitted two voice recordings used in audio preprocessing and voice model training."],
            ["Yvette Uwimpaye", "Submitted one image used in image preprocessing and face model training."],
        ],
    )

    document.add_heading("9. Limitations and Conclusion", level=1)
    document.add_paragraph(
        "The system demonstrates the complete multimodal sequence, including approved and denied paths. The merge "
        "was validated, the required feature CSV files were created, and every model reports accuracy, macro F1, "
        "and log loss."
    )
    document.add_paragraph(
        "The main limitation is media quantity. At the time of processing, Hassan had not uploaded media, Mahlet "
        "had no images, and Yvette had one image and no audio. The product data also has weak predictive signals. "
        "More original samples per member and more customer history would be needed for a stronger real system."
    )

    document.add_heading("10. Submission Links", level=1)
    add_bullets(
        document,
        [
            "GitHub repository: https://github.com/irachrist1/alu-machine_learning",
            "System demonstration video: Add the uploaded video URL before submission.",
        ],
    )

    document.add_heading("Data Sources", level=1)
    add_bullets(
        document,
        [
            "Customer social profiles: https://docs.google.com/spreadsheets/d/10up-WdC0a6egYaXLKiMQUotpOvaKZRZvYuWFQx4-RPQ/",
            "Customer transactions: https://docs.google.com/spreadsheets/d/1s4WOVm49lmLQ8d9QbbbdgAcRTNh3m0KCH_5ciiaRZw0/",
        ],
    )

    for new_section in document.sections[1:]:
        if new_section.start_type == WD_SECTION.NEW_PAGE:
            add_page_number(new_section.footer.paragraphs[0])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
